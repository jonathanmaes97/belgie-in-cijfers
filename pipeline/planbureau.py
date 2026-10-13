"""Inflatievooruitzichten en spilindex van het Federaal Planbureau, en de vooruitzichten van
andere instellingen voor België (verzameld door het Planbureau).

Bronnen (Excel, gepubliceerd door het Planbureau):
  https://www.plan.be/nl/data/indexcijfer-der-consumptieprijzen
  https://www.plan.be/nl/data/vooruitzichten-van-verschillende-instellingen-voor

Gemeten en verwachte maanden staan in hetzelfde bestand. Het Planbureau kleurt de
vooruitzichten oranje; die kleur lezen we uit. Lukt dat niet, dan geldt: alles vanaf
de maand waarin het bestand werd gepubliceerd is een vooruitzicht.

Spilindex: die is overschreden in de maand waarin de afgevlakte gezondheidsindex de
geldende spilindex bereikt of overschrijdt.
"""
from __future__ import annotations

import io
from datetime import date, datetime
from email.utils import parsedate_to_datetime

import requests

from .eurostat import BronFout

URL_INFLATIE = "https://www.plan.be/sites/default/files/documents/DATA_FOR_Inflation.xlsx"
URL_CONSENSUS = "https://www.plan.be/sites/default/files/documents/FOR_BE_NL.xlsx"
PAGINA_INFLATIE = "https://www.plan.be/nl/data/indexcijfer-der-consumptieprijzen"
PAGINA_CONSENSUS = "https://www.plan.be/nl/data/vooruitzichten-van-verschillende-instellingen-voor"
UA = {"User-Agent": "belgie-in-cijfers (open data dashboard; github.com/jonathanmaes97/belgie-in-cijfers)"}


def _getal(x):
    if x is None:
        return None
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).strip().replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def _is_oranje(cel) -> bool:
    try:
        kleur = cel.font.color
        rgb = (kleur.rgb or "") if kleur is not None and kleur.type == "rgb" else ""
    except Exception:
        return False
    if not isinstance(rgb, str) or len(rgb) < 6:
        return False
    r, g, b = int(rgb[-6:-4], 16), int(rgb[-4:-2], 16), int(rgb[-2:], 16)
    return r > 200 and 80 < g < 200 and b < 90


def lees_inflatie(inhoud: bytes, gepubliceerd: date) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(inhoud), data_only=True)
    ws = wb["NL"] if "NL" in wb.sheetnames else wb.worksheets[0]
    jaren, maanden = [], []
    deel = None
    kleur_gevonden = False
    for rij in ws.iter_rows():
        eerste = rij[0].value
        tekst = str(eerste or "").strip().lower()
        if tekst.startswith("jaarcijfers"):
            deel = "jaar"
            continue
        if tekst.startswith("maandcijfers"):
            deel = "maand"
            continue
        waarden = [c.value for c in rij[:7]] + [None] * 7
        if deel == "jaar" and isinstance(eerste, (int, float)) and 2000 < eerste < 2100:
            jaren.append({"jaar": int(eerste), "cpi": _getal(waarden[1]), "inflatie": _getal(waarden[2]),
                          "gezondheidsindex": _getal(waarden[3]), "gezondheid_groei": _getal(waarden[4])})
        elif deel == "maand" and isinstance(eerste, (datetime, date)):
            oranje = any(_is_oranje(c) for c in rij[:7])
            kleur_gevonden = kleur_gevonden or oranje
            maanden.append({"maand": eerste.strftime("%Y-%m"), "cpi": _getal(waarden[1]), "inflatie": _getal(waarden[2]),
                            "gezondheidsindex": _getal(waarden[3]), "gezondheid_groei": _getal(waarden[4]),
                            "afgevlakt": _getal(waarden[5]), "spilindex": _getal(waarden[6]), "_oranje": oranje})
    if not jaren or not maanden:
        raise BronFout("Inflatiebestand van het Planbureau heeft een onverwachte structuur.")
    grens = gepubliceerd.strftime("%Y-%m")
    for m in maanden:
        m["vooruitzicht"] = m.pop("_oranje") if kleur_gevonden else m["maand"] >= grens
    # overschrijdingen van de spilindex
    overschrijdingen = []
    for m in maanden:
        if m["afgevlakt"] is not None and m["spilindex"] is not None and m["afgevlakt"] >= m["spilindex"]:
            overschrijdingen.append({"maand": m["maand"], "vooruitzicht": m["vooruitzicht"], "spilindex": m["spilindex"]})
    laatste_gemeten = max((m["maand"] for m in maanden if not m["vooruitzicht"]), default=None)
    return {"jaren": jaren, "maanden": maanden, "overschrijdingen": overschrijdingen,
            "laatste_gemeten": laatste_gemeten, "kleur_uit_bestand": kleur_gevonden}


def lees_consensus(inhoud: bytes) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(inhoud), data_only=True, read_only=True)
    ws = wb.worksheets[0]
    rijen = list(ws.iter_rows(values_only=True))
    jaren = None
    for r in rijen:
        if r and r[1] is not None and isinstance(r[1], (int, float)) and 2000 < r[1] < 2100:
            jaren = [int(r[1]), int(r[2])]
            break
    if not jaren:
        raise BronFout("Consensusbestand van het Planbureau heeft een onverwachte structuur.")
    instellingen = []
    for r in rijen:
        if not r or not isinstance(r[0], str):
            continue
        cijfers = [_getal(v) for v in r[1:7]]
        if all(c is None for c in cijfers):
            continue
        upd = r[7] if len(r) > 7 else None
        instellingen.append({
            "naam": r[0].strip(),
            "groei": cijfers[0:2], "inflatie": cijfers[2:4], "saldo": cijfers[4:6],
            "update": upd.strftime("%Y-%m") if isinstance(upd, (datetime, date)) else (str(upd) if upd else None),
        })
    return {"jaren": jaren, "instellingen": instellingen}


def _haal(url: str) -> tuple[bytes, date]:
    try:
        r = requests.get(url, timeout=60, headers=UA)
    except requests.RequestException as e:
        raise BronFout(f"netwerkfout: {e}")
    if not r.ok:
        raise BronFout(f"HTTP {r.status_code} voor {url}")
    lm = r.headers.get("Last-Modified")
    gepubliceerd = parsedate_to_datetime(lm).date() if lm else date.today()
    return r.content, gepubliceerd


def haal_op() -> dict:
    inhoud, gepubliceerd = _haal(URL_INFLATIE)
    inflatie = lees_inflatie(inhoud, gepubliceerd)
    inhoud2, gepubliceerd2 = _haal(URL_CONSENSUS)
    consensus = lees_consensus(inhoud2)
    return {
        "inflatie": {**inflatie, "gepubliceerd": gepubliceerd.isoformat(), "bron": URL_INFLATIE, "pagina": PAGINA_INFLATIE},
        "consensus": {**consensus, "gepubliceerd": gepubliceerd2.isoformat(), "bron": URL_CONSENSUS, "pagina": PAGINA_CONSENSUS},
    }
