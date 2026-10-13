"""Wekelijkse prijs van mazout (gasolie voor verwarming) in België, uit de Weekly Oil Bulletin van de Europese Commissie.

Bron: https://energy.ec.europa.eu/data-and-analysis/weekly-oil-bulletin_en
Het historiekbestand bevat per week de consumentenprijs met alle belastingen, in euro per 1.000 liter.
Wij delen door 1.000 om de prijs per liter te tonen; verder wordt niets aangepast.
"""
from __future__ import annotations

import io
from datetime import date, datetime

import requests

from .eurostat import BronFout

PAGINA = "https://energy.ec.europa.eu/data-and-analysis/weekly-oil-bulletin_en"
HISTORIEK = ("https://energy.ec.europa.eu/document/download/906e60ca-8b6a-44e7-8589-652854d2fd3f_en"
             "?filename=Weekly_Oil_Bulletin_Prices_History_maticni_4web.xlsx")
BLAD = "Prices with taxes"
KOLOMMEN = {"BE": "BE_price_with_tax_heating_oil", "EU": "EU_price_with_tax_heating_oil"}
UA = {"User-Agent": "belgie-in-cijfers (open data dashboard; github.com/jonathanmaes97/belgie-in-cijfers)"}
WEKEN = 160  # ongeveer drie jaar


def lees(inhoud: bytes) -> dict[str, list[list]]:
    """Geeft per land een lijst [datum, prijs per liter, ""], oudste eerst. Stopt met een duidelijke fout als de opbouw wijzigt."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(inhoud), read_only=True, data_only=True)
    if BLAD not in wb.sheetnames:
        raise BronFout(f"Blad '{BLAD}' niet gevonden in de Weekly Oil Bulletin (wel: {wb.sheetnames}).")
    rijen = wb[BLAD].iter_rows(values_only=True)
    kop = list(next(rijen))
    next(rijen)                 # omschrijving van het product
    eenheden = list(next(rijen))  # eenheid per kolom
    pos = {}
    for land, naam in KOLOMMEN.items():
        if naam not in kop:
            raise BronFout(f"Kolom {naam} niet gevonden in de Weekly Oil Bulletin.")
        pos[land] = kop.index(naam)
        if str(eenheden[pos[land]]).strip() != "1000 l":
            raise BronFout(f"Onverwachte eenheid voor {naam}: {eenheden[pos[land]]!r} (verwacht: 1000 l).")
    uit: dict[str, list[list]] = {land: [] for land in KOLOMMEN}
    for rij in rijen:
        d = rij[0]
        if isinstance(d, datetime):
            d = d.date()
        if not isinstance(d, date):
            continue
        for land, i in pos.items():
            v = rij[i] if i < len(rij) else None
            if isinstance(v, (int, float)) and v > 0:
                uit[land].append([d.isoformat(), round(v / 1000, 4), ""])
    for land in uit:
        uit[land].sort(key=lambda p: p[0])
        uit[land] = uit[land][-WEKEN:]
    if len(uit["BE"]) < 10:
        raise BronFout("Te weinig weekprijzen voor België in de Weekly Oil Bulletin.")
    return uit


def haal_op() -> dict:
    try:
        r = requests.get(HISTORIEK, timeout=180, headers=UA)
    except requests.RequestException as e:
        raise BronFout(f"netwerkfout: {e}")
    if not r.ok:
        raise BronFout(f"HTTP {r.status_code} voor de Weekly Oil Bulletin.")
    try:
        reeksen = lees(r.content)
    except BronFout:
        raise
    except Exception as e:  # bv. geen geldig Excel-bestand meer
        raise BronFout(f"Weekly Oil Bulletin onleesbaar: {type(e).__name__}: {e}")
    return {"reeksen": [{
        "id": "mazout",
        "titel": {"nl": "Prijs van mazout"},
        "vraag": {"nl": "Wat kost een liter mazout?"},
        "uitleg": {"nl": "De gemiddelde prijs voor gasolie voor verwarming, alle belastingen inbegrepen, zoals België die elke week aan de Europese Commissie meldt. De Commissie geeft de prijs per 1.000 liter; wij delen door 1.000. Wat je zelf betaalt, hangt ook af van de hoeveelheid en je leverancier."},
        "eenheid": "€/l", "decimalen": 3,
        "punten": reeksen["BE"],
        "vergelijk": {"naam": {"nl": "EU"}, "punten": reeksen["EU"]},
        "api_url": HISTORIEK,
        "bron": {"instelling": "Europese Commissie", "naam": {"nl": "Europese Commissie (Weekly Oil Bulletin)"}, "url": PAGINA, "dataset": "Weekly Oil Bulletin"},
    }]}
